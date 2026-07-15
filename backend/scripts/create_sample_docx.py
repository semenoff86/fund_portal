"""Generate demo .docx templates with Jinja2 placeholders for docxtpl."""

import os
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
TEMPLATES_ROOT = Path(os.getenv("TEMPLATES_DIR", str(ROOT / "templates" / "documents")))
SOURCE_DIR = TEMPLATES_ROOT / "source"

TEMPLATES: dict[str, list[str]] = {
    "leave-unpaid.docx": [
        "СЛУЖЕБНАЯ ЗАПИСКА",
        "",
        "От {{ full_name }}",
        "Отдел: {{ department }}",
        "",
        "Прошу предоставить отпуск без сохранения заработной платы",
        "с {{ leave_start }} по {{ leave_end }} ({{ days_count }} календарных дней).",
        "",
        "Причина: {{ reason }}",
        "",
        "Дата: {{ document_date }}",
        "",
        "Подпись: ___________________ / {{ full_name }} /",
    ],
    "leave-paid.docx": [
        "СЛУЖЕБНАЯ ЗАПИСКА",
        "",
        "От {{ full_name }}",
        "Отдел: {{ department }}",
        "",
        "Прошу предоставить ежегодный оплачиваемый отпуск",
        "с {{ leave_start }} по {{ leave_end }} ({{ days_count }} календарных дней).",
        "",
        "{% if substitute_person %}На период отсутствия обязанности исполняет: {{ substitute_person }}.{% endif %}",
        "",
        "Дата: {{ document_date }}",
        "",
        "Подпись: ___________________ / {{ full_name }} /",
    ],
    "bonus.docx": [
        "СЛУЖЕБНАЯ ЗАПИСКА",
        "",
        "Прошу рассмотреть вопрос о выплате премии сотруднику {{ full_name }}",
        "(отдел: {{ department }}) за период: {{ bonus_period }}.",
        "",
        "Рекомендуемый размер премии: {{ bonus_amount }} руб.",
        "",
        "Обоснование:",
        "{{ achievement }}",
        "",
        "Дата: {{ document_date }}",
        "",
        "Подпись руководителя: ___________________",
    ],
    "business-trip.docx": [
        "СЛУЖЕБНАЯ ЗАПИСКА",
        "",
        "Прошу направить {{ full_name }} ({{ department }}) в служебную командировку.",
        "",
        "Место назначения: {{ destination }}",
        "Период: с {{ trip_start }} по {{ trip_end }}",
        "Цель: {{ trip_purpose }}",
        "{% if estimated_cost %}Предполагаемые расходы: {{ estimated_cost }} руб.{% endif %}",
        "",
        "Дата: {{ document_date }}",
        "",
        "Подпись: ___________________",
    ],
    "material-assistance.docx": [
        "СЛУЖЕБНАЯ ЗАПИСКА",
        "",
        "От {{ full_name }}",
        "Отдел: {{ department }}",
        "Телефон: {{ phone }}",
        "",
        "Прошу оказать материальную помощь в размере {{ assistance_amount }} руб.",
        "",
        "Основание:",
        "{{ reason }}",
        "",
        "Дата: {{ document_date }}",
        "",
        "Подпись: ___________________ / {{ full_name }} /",
    ],
}


def create_docx(filename: str, lines: list[str]) -> None:
    path = SOURCE_DIR / filename
    if path.exists():
        return
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    print(f"Created {path}")


def main():
    for filename, lines in TEMPLATES.items():
        create_docx(filename, lines)


if __name__ == "__main__":
    main()
